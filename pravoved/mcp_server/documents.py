import re
import uuid
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

REQUIRED = ["sender_name", "sender_address", "recipient", "recipient_address", "facts", "demand", "date"]

# kind -> (основания, предупреждение в конце)
CONSUMER_WARNING = (
    "В случае отказа или неответа я буду вынужден(а) обратиться в суд, а также в Роспотребнадзор. "
    "Одновременно прошу учесть, что при удовлетворении требований потребителя судом с продавца "
    "взыскивается штраф в размере 50% от суммы, присуждённой в пользу потребителя "
    "(п. 6 ст. 13 Закона РФ «О защите прав потребителей»)."
)

CLAIM_KINDS = {
    "consumer_refund": {
        "basis": [
            "Ст. 18 Закона РФ «О защите прав потребителей»: при обнаружении в товаре недостатков, "
            "не оговорённых продавцом, потребитель вправе отказаться от договора купли-продажи "
            "и потребовать возврата уплаченной за товар суммы, а также полного возмещения убытков.",
            "Ст. 22 Закона РФ «О защите прав потребителей»: требование о возврате уплаченной "
            "за товар суммы подлежит удовлетворению в течение десяти дней со дня предъявления требования.",
        ],
        "warning": CONSUMER_WARNING,
        "legal_basis": ["ЗоЗПП ст. 18", "ЗоЗПП ст. 22", "ЗоЗПП ст. 13 п. 6"],
    },
    "consumer_penalty": {
        "basis": [
            "Ст. 23 Закона РФ «О защите прав потребителей»: за нарушение сроков удовлетворения "
            "требований потребителя продавец уплачивает неустойку в размере одного процента цены товара "
            "за каждый день просрочки.",
            "Сумма неустойки не может превышать сумму, уплаченную потребителем по договору "
            "купли-продажи товара (с 01.02.2026).",
        ],
        "warning": CONSUMER_WARNING,
        "legal_basis": ["ЗоЗПП ст. 23", "ЗоЗПП ст. 13 п. 6"],
    },
    "salary_delay": {
        "basis": [
            "Ст. 236 Трудового кодекса РФ: при нарушении установленного срока выплаты заработной платы "
            "работодатель обязан выплатить её с процентами (денежной компенсацией) в размере не ниже "
            "одной сто пятидесятой действующей ключевой ставки Банка России от невыплаченных в срок сумм "
            "за каждый день задержки, начиная со дня, следующего за днём выплаты, по день фактического расчёта включительно.",
            "Ст. 142 Трудового кодекса РФ: при задержке выплаты заработной платы более чем на 15 дней "
            "работник вправе, известив работодателя в письменной форме, приостановить работу "
            "на весь период до выплаты задержанной суммы.",
        ],
        "warning": (
            "В случае невыполнения требований я буду вынужден(а) обратиться в государственную "
            "инспекцию труда, в прокуратуру и в суд."
        ),
        "legal_basis": ["ТК РФ ст. 236", "ТК РФ ст. 142"],
    },
}


def ics_escape(text: str) -> str:
    text = text.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,")
    return text.replace("\r\n", "\\n").replace("\n", "\\n")


def ics_fold(line: str) -> str:
    # строки длиннее 75 символов переносятся, продолжение начинается с пробела
    parts = [line[:75]]
    rest = line[75:]
    while rest:
        parts.append(" " + rest[:74])
        rest = rest[74:]
    return "\r\n".join(parts)


def parse_date(value) -> date:
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value))


def make_calendar_event(title, date, description="", remind_days_before=1, out_dir="data/out") -> dict:
    try:
        day = parse_date(date)
    except ValueError:
        return {"error": "дата должна быть в формате ГГГГ-ММ-ДД"}
    if remind_days_before < 0:
        return {"error": "remind_days_before не может быть отрицательным"}

    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//Pravoved//RU",
        "BEGIN:VEVENT",
        f"UID:{uuid.uuid4()}@pravoved",
        f"DTSTAMP:{stamp}",
        f"DTSTART;VALUE=DATE:{day:%Y%m%d}",
        f"DTEND;VALUE=DATE:{day + timedelta(days=1):%Y%m%d}",
        f"SUMMARY:{ics_escape(title)}",
    ]
    if description:
        lines.append(f"DESCRIPTION:{ics_escape(description)}")
    lines += [
        "BEGIN:VALARM",
        "ACTION:DISPLAY",
        f"DESCRIPTION:{ics_escape(title)}",
        f"TRIGGER:-P{remind_days_before}D",
        "END:VALARM",
        "END:VEVENT",
        "END:VCALENDAR",
    ]
    text = "\r\n".join(ics_fold(line) for line in lines) + "\r\n"

    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    path = out / f"event_{day:%Y%m%d}.ics"
    path.write_bytes(text.encode("utf-8"))
    return {"path": str(path), "date": day.isoformat()}


def draft_claim_docx(kind, data, out_dir="data/out") -> dict:
    if kind not in CLAIM_KINDS:
        return {"error": f"неизвестный вид претензии: {kind}", "kinds": list(CLAIM_KINDS)}
    missing = [k for k in REQUIRED if not str(data.get(k) or "").strip()]
    if missing:
        return {"error": "не хватает данных", "missing": missing}

    info = CLAIM_KINDS[kind]
    answer_days = data.get("answer_days") or 10
    # (текст, выравнивание, жирный)
    blocks = []

    right = [f"Кому: {data['recipient']}", f"Адрес: {data['recipient_address']}", ""]
    right.append(f"От кого: {data['sender_name']}")
    right.append(f"Адрес: {data['sender_address']}")
    if data.get("sender_phone"):
        right.append(f"Телефон: {data['sender_phone']}")
    for line in right:
        blocks.append((line, "right", False))

    blocks.append(("", "left", False))
    blocks.append(("ПРЕТЕНЗИЯ", "center", True))
    blocks.append(("", "left", False))
    blocks.append((data["facts"], "left", False))
    if data.get("amount"):
        blocks.append((f"Сумма, о которой идёт речь: {data['amount']}.", "left", False))
    if data.get("calc_breakdown"):
        blocks.append(("Расчёт:", "left", True))
        blocks.append((data["calc_breakdown"], "left", False))

    blocks.append(("Правовые основания:", "left", True))
    for text in info["basis"]:
        blocks.append((text, "left", False))

    blocks.append(("На основании изложенного ТРЕБУЮ:", "left", True))
    blocks.append((data["demand"], "left", False))
    blocks.append((f"Прошу исполнить требование в течение {answer_days} дней с момента получения настоящей претензии.", "left", False))
    blocks.append((info["warning"], "left", False))

    attachments = data.get("attachments") or []
    if attachments:
        blocks.append(("Приложения:", "left", True))
        for i, name in enumerate(attachments, 1):
            blocks.append((f"{i}. {name}", "left", False))

    blocks.append(("", "left", False))
    blocks.append((f"{data['date']}                    ______ / {data['sender_name']}", "left", False))

    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    align = {"right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}
    for text, where, bold in blocks:
        p = doc.add_paragraph()
        p.alignment = align[where]
        run = p.add_run(text)
        run.bold = bold

    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    safe_date = re.sub(r"[^0-9A-Za-zА-Яа-я_-]+", "-", str(data["date"])).strip("-")
    path = out / f"claim_{kind}_{safe_date}.docx"
    doc.save(path)

    preview = "\n".join(text for text, _, _ in blocks)
    return {"path": str(path), "preview": preview, "legal_basis": info["legal_basis"]}
